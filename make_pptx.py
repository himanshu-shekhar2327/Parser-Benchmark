"""
Rebuilds ONLY data/raw/invoice.pptx so the slide really contains every fact
the checker expects. (The first version left the invoice number, VAT number,
date and totals off the slide, so those 'misses' were the test's fault.)
"""

from pptx import Presentation
from pptx.util import Inches

LINES = [("USB cable", 5, 120.00), ("Desk lamp", 2, 350.00)]
SUBTOTAL, VAT, TOTAL = 1300.00, 195.00, 1495.00

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[5])        # title-only layout
slide.shapes.title.text = "TAX INVOICE - Example Supplies (Pty) Ltd"

# Header facts as a text box under the title
head = slide.shapes.add_textbox(Inches(0.5), Inches(1.7), Inches(9), Inches(0.6))
head.text_frame.text = "VAT No: 4000000001   Invoice No: INV-0001   Date: 05/10/2026"

# Line-item table
tbl = slide.shapes.add_table(len(LINES) + 1, 4, Inches(0.5), Inches(2.5),
                             Inches(9), Inches(1.8)).table
for i, c in enumerate(["Description", "Qty", "Unit Price", "Amount"]):
    tbl.cell(0, i).text = c
for r, (desc, q, p) in enumerate(LINES, start=1):
    for i, v in enumerate([desc, str(q), f"{p:.2f}", f"{q * p:.2f}"]):
        tbl.cell(r, i).text = v

# Totals as a text box under the table
foot = slide.shapes.add_textbox(Inches(0.5), Inches(4.7), Inches(9), Inches(0.6))
foot.text_frame.text = f"Subtotal: {SUBTOTAL:.2f}   VAT 15%: {VAT:.2f}   Total: {TOTAL:.2f}"

# Speaker notes: does the parser read them?
slide.notes_slide.notes_text_frame.text = "NOTES-MARKER: speaker note text"
prs.save("data/raw/invoice.pptx")
print("Rebuilt data/raw/invoice.pptx")
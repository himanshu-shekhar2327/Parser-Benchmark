"""Generates one small test invoice per format into data/raw/ (fake data)."""

from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation

RAW = Path("data/raw")
RAW.mkdir(parents=True, exist_ok=True)

# Fake invoice. The arithmetic is internally consistent on purpose, so a
# correct parse should pass the "lines and totals add up" gate.
HEADER = ["TAX INVOICE", "Example Supplies (Pty) Ltd", "VAT No: 4000000001",
          "Invoice No: INV-0001", "Date: 05/10/2026"]
LINES = [("USB cable", 5, 120.00), ("Desk lamp", 2, 350.00)]   # 600.00 + 700.00
SUBTOTAL = sum(q * p for _, q, p in LINES)                     # 1300.00
VAT = round(SUBTOTAL * 0.15, 2)                                # 195.00
TOTAL = SUBTOTAL + VAT                                         # 1495.00

# ---- DOCX ----
d = Document()
d.add_heading(HEADER[0], 1)
for h in HEADER[1:]:
    d.add_paragraph(h)
t = d.add_table(rows=1, cols=4)
for i, c in enumerate(["Description", "Qty", "Unit Price", "Amount"]):
    t.rows[0].cells[i].text = c
for desc, q, p in LINES:
    row = t.add_row().cells
    row[0].text, row[1].text, row[2].text, row[3].text = desc, str(q), f"{p:.2f}", f"{q*p:.2f}"
d.add_paragraph(f"Subtotal: {SUBTOTAL:.2f}   VAT 15%: {VAT:.2f}   Total: {TOTAL:.2f}")
d.save(RAW / "invoice.docx")

# ---- XLSX (merged title cell + a hidden sheet) ----
wb = Workbook()
ws = wb.active
ws.title = "Invoice"
ws.merge_cells("A1:D1")
ws["A1"] = HEADER[0]
for i, h in enumerate(HEADER[1:], start=2):
    ws[f"A{i}"] = h
r0 = len(HEADER) + 2
ws.append([])
ws.append(["Description", "Qty", "Unit Price", "Amount"])
for desc, q, p in LINES:
    ws.append([desc, q, p, q * p])
ws.append(["", "", "Subtotal", SUBTOTAL])
ws.append(["", "", "VAT 15%", VAT])
ws.append(["", "", "Total", TOTAL])
hidden = wb.create_sheet("HiddenNotes")
hidden["A1"] = "HIDDEN-SHEET-MARKER"      # does the parser see hidden sheets?
hidden.sheet_state = "hidden"
wb.save(RAW / "invoice.xlsx")

# ---- PPTX (table + speaker notes) ----
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[5])   # title-only layout
slide.shapes.title.text = HEADER[0] + " - " + HEADER[1]
from pptx.util import Inches
tbl = slide.shapes.add_table(len(LINES) + 1, 4, Inches(0.5), Inches(2), Inches(9), Inches(2)).table
for i, c in enumerate(["Description", "Qty", "Unit Price", "Amount"]):
    tbl.cell(0, i).text = c
for r, (desc, q, p) in enumerate(LINES, start=1):
    for i, v in enumerate([desc, str(q), f"{p:.2f}", f"{q*p:.2f}"]):
        tbl.cell(r, i).text = v
slide.notes_slide.notes_text_frame.text = "NOTES-MARKER: speaker note text"
prs.save(RAW / "invoice.pptx")

# ---- HTML ----
rows = "".join(f"<tr><td>{d_}</td><td>{q}</td><td>{p:.2f}</td><td>{q*p:.2f}</td></tr>"
               for d_, q, p in LINES)
html = (f"<html><body><h1>{HEADER[0]}</h1>" + "".join(f"<p>{h}</p>" for h in HEADER[1:]) +
        f"<table border=1><tr><th>Description</th><th>Qty</th><th>Unit Price</th>"
        f"<th>Amount</th></tr>{rows}</table>"
        f"<p>Subtotal: {SUBTOTAL:.2f} VAT 15%: {VAT:.2f} Total: {TOTAL:.2f}</p></body></html>")
(RAW / "invoice.html").write_text(html, encoding="utf-8")

# ---- PNG: clean, then degraded (tilt + noise + blur) ----
try:
    font = ImageFont.truetype("arial.ttf", 30)       # exists on Windows
except OSError:
    font = ImageFont.load_default()
img = Image.new("L", (1000, 700), 255)
draw = ImageDraw.Draw(img)
text = HEADER + [""] + ["Description        Qty   Unit Price   Amount"] + \
       [f"{d_:<18}{q:<6}{p:<13.2f}{q*p:.2f}" for d_, q, p in LINES] + \
       ["", f"Subtotal: {SUBTOTAL:.2f}   VAT 15%: {VAT:.2f}   Total: {TOTAL:.2f}"]
for i, line in enumerate(text):
    draw.text((40, 30 + i * 50), line, fill=0, font=font)
img.save(RAW / "invoice_clean.png")

noisy = Image.blend(img, Image.effect_noise(img.size, 40), 0.25)
tilted = noisy.rotate(3, expand=True, fillcolor=255)
tilted.filter(ImageFilter.GaussianBlur(1.3)).save(RAW / "invoice_degraded.png")

print("Created samples in", RAW)